"""Build a Word validation report containing evidence for every test_img file."""
from __future__ import annotations

import io
import json
from collections import Counter
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
RECOGNITION = ROOT / "artifacts" / "full_image_validation" / "recognition_evidence.json"
SCORED = ROOT / "artifacts" / "full_image_validation" / "scored_evidence.json"
OUT = ROOT / "artifacts" / "會計憑證辨識完整驗證報告_40張樣本.docx"

FIELD_LABELS = {
    "doc_type_key": "單據類型", "voucher_id": "憑證號碼", "date": "日期",
    "seller_tax_id": "賣方統編", "buyer_tax_id": "買方統編", "subtotal": "未稅金額",
    "tax_amount": "營業稅額", "total_amount": "總計", "debit_total": "借方合計",
    "credit_total": "貸方合計",
}
TYPE_LABELS = {"invoice": "統一發票", "receipt": "收據", "journal_voucher": "會計傳票", "other": "其他"}


def set_font(run, size=10, bold=False, colour=None):
    run.font.name = "Microsoft JhengHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    if colour:
        run.font.color.rgb = RGBColor(*colour)


def shade(cell, fill):
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(node)


def borders(table, colour="D9D9D9"):
    el = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), colour)
        el.append(tag)
    table._tbl.tblPr.append(el)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_text(cell, value, size=8.8, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, colour=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(value if value not in (None, "") else "—"))
    set_font(r, size, bold, colour)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def add_table(doc, headers, rows, widths, font_size=8.8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    borders(table)
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = widths[idx]
        shade(cell, "1F4E78")
        add_text(cell, text, font_size, True, WD_ALIGN_PARAGRAPH.CENTER, (255, 255, 255))
    set_repeat_header(table.rows[0])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].width = widths[idx]
            if row_index % 2:
                shade(cells[idx], "F4F8FB")
            alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in (0, len(row) - 1) else WD_ALIGN_PARAGRAPH.LEFT
            add_text(cells[idx], value, font_size, False, alignment)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1, page_break_before=False):
    p = doc.add_paragraph(style=f"Heading {level}")
    # Word's built-in heading style can carry a page-break-before setting.
    # Explicitly clear it so the deliberate page breaks below do not create blanks.
    p.paragraph_format.page_break_before = page_break_before
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6 if level == 1 else 4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_font(r, 15 if level == 1 else 11.5, True)
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_font(r, 9.8)
    return p


def picture(run, path: Path, width: float, alt_text: str, max_height: float | None = None):
    with Image.open(path) as source:
        source_width, source_height = source.size
        buf = io.BytesIO()
        source.convert("RGB").save(buf, format="PNG")
    buf.seek(0)
    display_width = width
    display_height = width * source_height / source_width
    if max_height is not None and display_height > max_height:
        display_height = max_height
        display_width = display_height * source_width / source_height
    inline = run.add_picture(buf, width=Inches(display_width), height=Inches(display_height))
    inline._inline.docPr.set("descr", alt_text)


def add_image_cell(cell, path: Path, width: float, caption: str, alt_text: str, max_height: float | None = None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    picture(p.add_run(), path, width, alt_text, max_height)
    c = cell.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(0)
    r = c.add_run(caption)
    set_font(r, 7.8, False, (80, 80, 80))
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    set_cell_margins(cell, top=40, start=40, bottom=40, end=40)


def add_image_pair(doc, left_path, left_caption, left_alt, left_width, right_path, right_caption, right_alt, right_width,
                   left_max_height=None, right_max_height=None):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    borders(table, "E5E7EB")
    table.columns[0].width = Inches(4.0)
    table.columns[1].width = Inches(6.2)
    add_image_cell(table.cell(0, 0), left_path, left_width, left_caption, left_alt, left_max_height)
    add_image_cell(table.cell(0, 1), right_path, right_width, right_caption, right_alt, right_max_height)
    return table


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11), Inches(8.5)
    section.top_margin, section.bottom_margin = Inches(0.38), Inches(0.35)
    section.left_margin, section.right_margin = Inches(0.36), Inches(0.36)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft JhengHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    normal.font.size = Pt(9.8)
    for name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Microsoft JhengHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        style.font.color.rgb = RGBColor(0, 0, 0)
    title_style_pr = doc.styles["Title"]._element.get_or_add_pPr()
    title_style_border = title_style_pr.find(qn("w:pBdr"))
    if title_style_border is not None:
        title_style_pr.remove(title_style_border)
    doc.styles["Heading 1"].paragraph_format.page_break_before = False
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("會計憑證辨識系統 全資料集驗證報告"), 8, False, (105, 105, 105))
    return doc


def fmt(value):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return "—" if value in (None, "") else str(value)


def main():
    recognition = json.loads(RECOGNITION.read_text(encoding="utf-8"))
    scored = json.loads(SCORED.read_text(encoding="utf-8"))
    scored_by_name = {Path(x["source_image"]).name: x for x in scored["results"]}
    records = sorted(recognition["results"], key=lambda x: Path(x["source_image"]).name)
    correct, total, accuracy = scored["correct_fields"], scored["total_scored_fields"], scored["field_accuracy"]
    type_counter = Counter(x.get("doc_type_key") or "error" for x in records)
    failed = [x for x in records if x.get("error")]
    doc = setup_document()

    title = doc.add_paragraph(style="Title")
    # Remove the default Word Title bottom rule; the report uses its own clean layout.
    p_pr = title._p.get_or_add_pPr()
    title_border = p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        p_pr.remove(title_border)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(5)
    set_font(title.add_run("會計憑證辨識完整驗證報告"), 22, True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(9)
    set_font(subtitle.add_run(f"test_img 全部 40 張樣本  報告日期 {date.today().isoformat()}"), 10.5, False, (75, 75, 75))

    add_heading(doc, "驗證結論")
    add_body(doc, f"本次以 test_img 的 40 張原始憑證逐張執行目前系統辨識流程，並由人工逐張核對可清楚辨讀的核心欄位。共納入 {total} 個人工基準欄位，辨識正確 {correct} 個，欄位準確率為 {accuracy:.1%}。所有圖片均保留原始圖、辨識標註圖、欄位結果及明細結果。")
    add_table(doc, ["指標", "結果", "說明"], [
        ["測試範圍", "40 張", "ref/test_img/test_001.jpg 至 test_040.jpg"],
        ["辨識完成", f"{len(records) - len(failed)} 張", "以目前系統設定逐張執行；失敗筆數另列"],
        ["可計分欄位", f"{total} 項", "僅納入原始影像中可人工明確讀取的核心欄位"],
        ["正確欄位", f"{correct} 項", "採欄位完全一致比對"],
        ["欄位準確率", f"{accuracy:.1%}", f"{correct} / {total}"],
        ["辨識引擎", "VLM 直讀備援", "本機 OCR 未完成穩定輸出，未更改系統保存設定"],
    ], [Cm(3.3), Cm(3.2), Cm(19.0)], 9.0)

    add_heading(doc, "資料分布與評分方法")
    add_table(doc, ["實際判定類型", "樣本數", "評分欄位"], [
        [TYPE_LABELS.get(key, key), value, "類型、憑證號碼、日期、統編及金額欄位，依原圖可見性納入" ]
        for key, value in sorted(type_counter.items())
    ], [Cm(5.0), Cm(3.0), Cm(17.5)], 8.8)
    add_body(doc, "比對前會統一移除空白、逗號與大小寫差異，數字 2 與 2.0 視為一致；日期均以西元 YYYY-MM-DD 比對。影像未清楚呈現、或無法自原圖可靠確認的欄位不計入分母。")

    add_heading(doc, "分欄位準確率", page_break_before=True)
    by_field_rows = []
    for key, values in sorted(scored["per_field"].items(), key=lambda x: (-x[1]["total"], x[0])):
        pct = values["correct"] / values["total"] if values["total"] else 0
        by_field_rows.append([FIELD_LABELS.get(key, key), f"{values['correct']} / {values['total']}", f"{pct:.1%}"])
    add_table(doc, ["欄位", "正確 / 計分", "準確率"], by_field_rows, [Cm(7.0), Cm(6.0), Cm(5.0)], 9.0)

    add_heading(doc, "逐張彙總")
    summary_rows = []
    for record in records:
        name = Path(record["source_image"]).name
        result = scored_by_name.get(name, {})
        got, den = result.get("correct_fields", 0), result.get("total_scored_fields", 0)
        pct = f"{got / den:.0%}" if den else "—"
        wrong = "、".join(FIELD_LABELS.get(x["field"], x["field"]) for x in result.get("comparisons", []) if not x["match"])
        summary_rows.append([name, TYPE_LABELS.get(record.get("doc_type_key"), record.get("doc_type_key", "辨識失敗")), f"{got}/{den}", pct, wrong or "—"])
    add_table(doc, ["樣本", "辨識類型", "正確 / 計分", "準確率", "不符欄位"], summary_rows,
              [Cm(3.2), Cm(4.2), Cm(3.1), Cm(2.4), Cm(12.1)], 7.9)

    add_heading(doc, "辨識結果分析", page_break_before=True)
    mismatch = Counter()
    for result in scored["results"]:
        for item in result["comparisons"]:
            if not item["match"]:
                mismatch[item["field"]] += 1
    analysis_rows = [[FIELD_LABELS.get(k, k), v, "需優先加入格式檢核與人工覆核規則"] for k, v in mismatch.most_common()]
    if analysis_rows:
        add_table(doc, ["主要不符欄位", "不符筆數", "建議"], analysis_rows, [Cm(6.0), Cm(4.0), Cm(15.0)], 8.9)
    add_body(doc, "手寫發票與收據的欄位位置、筆跡及印章重疊差異最大，常影響憑證號碼、日期、統編與金額讀取。會計傳票屬密集多欄表格，現行 VLM 直讀備援未能穩定提供定位方框，應先修復 OCR，再以 OCR 座標輔助欄位與借貸欄判讀。")
    add_heading(doc, "後續改善建議")
    add_table(doc, ["優先項目", "目的"], [
        ["修復 OCR 並重新測試", "恢復文字座標與欄位框精修，特別改善密集傳票與印章內統編"],
        ["憑證號碼與統編格式檢核", "限制統編為 8 位數及發票號碼格式，攔截缺字、漏字或將電話誤寫入統編的結果"],
        ["金額交叉驗算", "對應稅發票檢查未稅金額加稅額是否等於總計；明細金額與總計應一致"],
        ["建立審核回寫資料", "將人工修正後的欄位保存為長期標準答案，支援後續模型與規則的回歸驗證"],
    ], [Cm(6.0), Cm(19.0)], 9.0)

    # Each sample occupies two intentional landscape pages. This preserves the
    # readability of portrait receipts while keeping every required evidence
    # view (raw, annotated, fields and details) together with its sample.
    for idx, record in enumerate(records, 1):
        name = Path(record["source_image"]).name
        result = scored_by_name.get(name, {})
        got, den = result.get("correct_fields", 0), result.get("total_scored_fields", 0)
        type_label = TYPE_LABELS.get(record.get("doc_type_key"), record.get("doc_type_key", "辨識失敗"))
        add_heading(doc, f"樣本 {idx:02d}  {name}", page_break_before=True)
        subtitle = doc.add_paragraph()
        subtitle.paragraph_format.keep_with_next = True
        subtitle.paragraph_format.space_after = Pt(3)
        if record.get("error"):
            line = f"辨識失敗：{record['error']}"
        else:
            line = f"辨識類型：{type_label}　人工核心欄位比對：{got} / {den} 正確"
        set_font(subtitle.add_run(line), 9.4, False, (60, 60, 60))
        raw = ROOT / record["source_image"]
        visuals = record.get("visuals") or {}
        if record.get("error") or not visuals:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture(p.add_run(), raw, 5.4, f"{name} 原始測試圖片")
            continue
        add_image_pair(doc, raw, "原始圖片", f"{name} 原始測試圖片", 3.85,
                       ROOT / visuals["annotated"], "辨識後圖片：數字與欄位方框", f"{name} 辨識標註圖", 5.85,
                       left_max_height=5.55, right_max_height=5.55)
        if record.get("box_count", 0) == 0:
            note = doc.add_paragraph()
            note.paragraph_format.space_before = Pt(1)
            note.paragraph_format.space_after = Pt(0)
            set_font(note.add_run("註：此密集表單未取得可用欄位座標，標註圖保留原始影像；欄位與明細結果仍為本次系統輸出。"), 7.6, False, (105, 105, 105))
        add_heading(doc, "欄位與明細結果", page_break_before=True)
        fields_subtitle = doc.add_paragraph()
        fields_subtitle.paragraph_format.keep_with_next = True
        fields_subtitle.paragraph_format.space_after = Pt(3)
        set_font(fields_subtitle.add_run(f"樣本 {idx:02d}　{name}｜右側欄位與明細畫面均為本次系統輸出截圖。"), 9.0, False, (60, 60, 60))
        add_image_pair(doc, ROOT / visuals["fields"], "右側欄位辨識結果", f"{name} 欄位結果截圖", 4.9,
                       ROOT / visuals["details"], "明細辨識結果", f"{name} 明細結果截圖", 4.9,
                       left_max_height=5.55, right_max_height=5.55)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
