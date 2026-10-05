"""Build the initial two-image recognition validation report as a DOCX."""
from __future__ import annotations

import json
import io
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
EVIDENCE = ROOT / "artifacts" / "two_image_validation" / "validation_evidence.json"
OUT = ROOT / "artifacts" / "會計憑證辨識初步驗證報告_兩張樣本.docx"
ATTACHED_BOXES = Path(r"C:\Users\itri\AppData\Local\Temp\codex-clipboard-485ea303-56bd-4861-9893-2657a73b6504.png")
ATTACHED_FIELDS = Path(r"C:\Users\itri\AppData\Local\Temp\codex-clipboard-d49d076b-6b48-4cb0-b6bd-a0454d05d5fa.png")
ATTACHED_DETAILS = Path(r"C:\Users\itri\AppData\Local\Temp\codex-clipboard-983e02de-c02c-4536-9259-421d7c283258.png")

LABELS = {
    "voucher_id": "發票號碼",
    "date": "發票日期",
    "invoice_type": "發票格式",
    "tax_type": "課稅別",
    "buyer_tax_id": "買方統編",
    "item_description": "品名",
    "quantity": "數量",
    "unit_price": "單價",
    "item_amount": "品項金額",
    "subtotal": "未稅銷售額",
    "tax_amount": "營業稅額",
    "total_amount": "總計含稅",
}


def set_font(run, size=10.5, bold=False, color=None):
    run.font.name = "Microsoft JhengHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:fill"), fill)
    tc_pr.append(el)


def borders(table):
    tbl_pr = table._tbl.tblPr
    node = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), "D9D9D9")
        node.append(tag)
    tbl_pr.append(node)


def cell_text(cell, text, *, bold=False, color=None, size=9.5, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text if text not in (None, "") else "—"))
    set_font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_font(r, size=15 if level == 1 else 12, bold=True)
    return p


def add_body(doc, text, *, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_table(doc, headers, rows, widths=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    borders(table)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, "1F4E78")
        cell_text(cell, h, bold=True, color=(255, 255, 255), size=font_size,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        if widths:
            cell.width = widths[i]
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    header_marker = OxmlElement("w:tblHeader")
    header_marker.set(qn("w:val"), "true")
    header_pr.append(header_marker)
    for r_i, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if r_i % 2 == 1:
                shade(cells[i], "F4F8FB")
            align = WD_ALIGN_PARAGRAPH.CENTER if i >= len(row) - 1 else WD_ALIGN_PARAGRAPH.LEFT
            cell_text(cells[i], value, size=font_size, align=align)
            if widths:
                cells[i].width = widths[i]
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def format_value(value):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (int, float)):
        return f"{value:,.0f}"
    return value if value not in (None, "") else "—"


def add_source_image(run, path, alt_text, width_inches=5.85):
    """Embed a lossless in-memory PNG so all source JPEG variants render in Word."""
    with Image.open(path) as image:
        buf = io.BytesIO()
        image.convert("RGB").save(buf, format="PNG")
    buf.seek(0)
    shape = run.add_picture(buf, width=Inches(width_inches))
    shape._inline.docPr.set("descr", alt_text)


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.70)
    section.right_margin = Inches(0.70)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft JhengHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    normal.font.size = Pt(10.5)
    for name in ("Title", "Heading 1", "Heading 2"):
        style = styles[name]
        style.font.name = "Microsoft JhengHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        style.font.color.rgb = RGBColor(0, 0, 0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("會計憑證辨識系統 初步驗證報告")
    set_font(run, size=8.5, color=(100, 100, 100))
    return doc


def main():
    evidence = json.loads(EVIDENCE.read_text(encoding="utf-8"))
    results = evidence["results"]
    correct = evidence["correct_fields"]
    total = evidence["total_scored_fields"]
    accuracy = evidence["field_accuracy"]
    mismatches = [
        (Path(result["source_image"].replace("\\", "/")).name, key, comparison)
        for result in results
        for key, comparison in result["comparisons"].items()
        if not comparison["match"]
    ]
    voucher_hits = sum(
        result["comparisons"]["voucher_id"]["match"] for result in results
    )
    sample_breakdown = "；".join(
        f"{Path(result['source_image'].replace('\\', '/')).name} {result['correct_fields']}/{result['total_scored_fields']}"
        for result in results
    )
    doc = setup_document()

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("會計憑證辨識初步驗證報告")
    set_font(run, size=22, bold=True)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(14)
    run = sub.add_run(f"兩張樣本測試  報告日期 {date.today().isoformat()}")
    set_font(run, size=11, color=(80, 80, 80))

    add_heading(doc, "驗證結論")
    add_body(doc, f"本次先以 test_001.jpg 與 test_002.jpg 進行初步驗證。以人工目視確認的 {total} 項欄位為基準，系統辨識正確 {correct} 項，欄位準確率為 {accuracy:.1%}。不符項目涵蓋日期、欄位值格式、發票號碼與買方統編，金額欄位則全部正確。")
    add_table(doc, ["指標", "結果", "說明"], [
        ["測試圖片", "2 張", "test_001.jpg、test_002.jpg"],
        ["可計分欄位", f"{total} 項", "只納入影像中可明確人工判讀的欄位"],
        ["正確欄位", f"{correct} 項", sample_breakdown],
        ["欄位準確率", f"{accuracy:.1%}", f"{correct} / {total}"],
        ["金額欄位", "12 / 12", "數量、單價、品項金額、未稅、稅額、總計均正確"],
        ["發票號碼", f"{voucher_hits} / {len(results)}", "採完整英文字首加數字的完全一致比對"],
    ], widths=[Cm(3.2), Cm(2.7), Cm(11.4)])

    add_heading(doc, "測試範圍與方法")
    add_body(doc, "測試資料取自 ref/test_img，原始影像未改寫。辨識前先複製到測試工作區，避免自動轉向功能影響原檔。")
    add_table(doc, ["項目", "本次作法"], [
        ["資料範圍", "僅 test_001.jpg 與 test_002.jpg，兩者均為三聯式統一發票"],
        ["人工基準", "由原始影像逐欄轉錄；未能清楚確認的欄位不計入準確率"],
        ["模型設定", "OCR：paddle/pp-ocrv5；結構辨識：ollama/qwen2.5vl:7b；科目分類：ollama/qwen2.5:7b"],
        ["計分規則", "移除空白與千分位逗號後採完全一致比對；數值 2 與 2.0 視為一致"],
        ["執行限制", "本機 OCR 執行在載入模型後未完成輸出，本報告的有效辨識紀錄為 VLM 直讀備援結果。因此本次數字不代表已完成 OCR 加 VLM 全流程驗收。"],
    ], widths=[Cm(3.2), Cm(14.1)], font_size=9.1)

    for index, result in enumerate(results, start=1):
        doc.add_page_break()
        source_name = Path(result["source_image"]).name
        recognized = result["recognized"]
        add_heading(doc, f"樣本 {index}  {source_name}")
        add_body(doc, f"判定類型：{recognized.get('doc_type', '—')}。本樣本 {result['correct_fields']} / {result['total_scored_fields']} 項計分欄位正確。")
        img_p = ROOT / result["source_image"].replace("\\", "/")
        image_paragraph = doc.add_paragraph()
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.paragraph_format.space_after = Pt(3)
        add_source_image(image_paragraph.add_run(), img_p, f"{source_name} 原始測試圖片")
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(7)
        r = cap.add_run("圖 1 原始測試圖片" if index == 1 else "圖 2 原始測試圖片")
        set_font(r, size=8.8, color=(80, 80, 80))

        rows = []
        for key, comparison in result["comparisons"].items():
            rows.append([
                LABELS[key],
                format_value(comparison["expected"]),
                format_value(comparison["observed"]),
                "正確" if comparison["match"] else "不符",
            ])
        add_table(doc, ["欄位", "人工基準", "系統辨識結果", "判定"], rows,
                  widths=[Cm(3.1), Cm(4.4), Cm(6.2), Cm(2.0)], font_size=8.8)
        fields = recognized.get("fields") or {}
        add_body(doc, f"其他辨識明細：賣方名稱「{fields.get('seller_name') or '—'}」；賣方統編「{fields.get('seller_tax_id') or '—'}」；模型信心標示為「{recognized.get('confidence') or '—'}」。本次科目分類未產生分錄，且未提供人工核定分錄，因此不納入會計科目準確率。")

    # User-provided UI evidence is tied to test_011.jpg (the source image is
    # visually identical to the background in the annotated screenshot).
    doc.add_page_break()
    add_heading(doc, "系統辨識畫面證據 test_011.jpg")
    add_body(doc, "以下截圖補充展示使用者在系統畫面中看到的完整證據鏈：原始圖片、帶數字與方框的辨識圖片、右側欄位結果，以及明細結果。此畫面樣本獨立於前述兩張圖片的 24 欄位準確率計算。")
    add_heading(doc, "原始圖片與辨識方框", level=2)
    raw = ROOT / "ref" / "test_img" / "test_011.jpg"
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(2)
    add_source_image(p.add_run(), raw, "test_011.jpg 原始圖片", width_inches=3.0)
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after = Pt(3)
    set_font(c.add_run("圖 3 原始圖片 test_011.jpg"), size=8.8, color=(80, 80, 80))
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(2)
    add_source_image(p.add_run(), ATTACHED_BOXES, "test_011.jpg 辨識後圖片，包含數字與欄位方框", width_inches=5.85)
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after = Pt(4)
    set_font(c.add_run("圖 4 辨識後圖片：數字與彩色方框"), size=8.8, color=(80, 80, 80))

    doc.add_page_break()
    add_heading(doc, "欄位與明細結果截圖", level=2)
    add_body(doc, "右側欄位畫面顯示系統實際填入的欄位值、信心度與檢核提示；下方明細畫面顯示品名、數量、單價與合計。")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(2)
    add_source_image(p.add_run(), ATTACHED_FIELDS, "test_011.jpg 右側欄位辨識結果", width_inches=3.35)
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after = Pt(4)
    set_font(c.add_run("圖 5 右側欄位辨識結果"), size=8.8, color=(80, 80, 80))
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(2)
    add_source_image(p.add_run(), ATTACHED_DETAILS, "test_011.jpg 明細辨識結果", width_inches=6.2)
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after = Pt(4)
    set_font(c.add_run("圖 6 明細辨識結果"), size=8.8, color=(80, 80, 80))

    doc.add_page_break()
    add_heading(doc, "結果分析")
    add_body(doc, "本次兩張樣本均可辨識為統一發票，金額與稅額的算術關係也正確。嚴格的欄位比對仍找出日期、欄位值格式與識別碼的錯誤。完整發票號碼與買方統編是去重、對帳與稅務勾稽的重要鍵值，應優先建立格式檢核。")
    impact = {
        "voucher_id": "完整憑證識別碼不一致，影響去重與對帳",
        "date": "入帳日期可能誤置年度或期間",
        "invoice_type": "欄位未依契約輸出標準值，影響後續規則判斷",
        "buyer_tax_id": "買方識別碼缺漏，影響統編勾稽",
    }
    mismatch_rows = [
        [sample, LABELS[key], format_value(item["expected"]), format_value(item["observed"]), impact.get(key, "需人工複核")]
        for sample, key, item in mismatches
    ]
    add_table(doc, ["錯誤樣本", "欄位", "人工基準", "系統結果", "影響"], mismatch_rows,
              widths=[Cm(2.6), Cm(2.6), Cm(3.5), Cm(3.3), Cm(4.0)], font_size=8.9)

    add_heading(doc, "準確率解讀與後續驗證")
    add_body(doc, f"{accuracy:.1%} 僅描述這兩張三聯式發票、這 {total} 個已人工確認欄位的初步表現，不能外推為整個資料集或所有憑證類型的系統準確率。尤其目前尚未取得科目分類的人工標準答案，也尚未驗證 OCR 加 VLM 的完整串接結果。另一次重跑曾出現不同欄位結果，表示還需納入重複執行的穩定性量測。")
    add_table(doc, ["建議項目", "目的"], [
        ["修復並重跑 OCR", "確認 OCR 是否可正常完成，並以同兩張圖片比較 OCR 加 VLM 與 VLM 直讀的差異"],
        ["補齊 40 張人工標註", "針對每張圖片建立欄位標準答案，才能計算整體與分欄位的準確率"],
        ["分憑證類型統計", "將統一發票、收據、繳款書與其他類型分開報告，避免單一類型掩蓋差異"],
        ["新增憑證號碼檢核", "以格式規則檢查英文字首加 8 位數字，降低完整識別碼遺漏的風險"],
    ], widths=[Cm(4.0), Cm(11.7)], font_size=9.1)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
